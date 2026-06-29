#!/usr/bin/env python3
"""Gera planilha Excel editável do fluxograma de investimentos."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
LOGO_PNG = BASE / "logo-3m-agropecuaria.png"
OUTPUT = BASE / "fluxograma-investimentos.xlsx"

NAVY = "003876"
GREEN = "0B610B"
WHITE = "FFFFFF"
BG_1 = "EEF2F7"
BG_2 = "EDF5EE"
BG_ALT = "F8FAFC"
BORDER = "CBD5E1"

SECTIONS = [
    ("Investimentos em Andamento", NAVY, BG_1),
    ("Investimentos Previstos com Orçamento", GREEN, BG_2),
    ("Investimentos Previstos sem Orçamento", NAVY, "FFFFFF"),
]

COLUMNS = ["Valor", "Status", "Prioridade"]
INITIAL_ROWS = 4


def thin_border():
    side = Side(style="thin", color=BORDER)
    return Border(left=side, right=side, top=side, bottom=side)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Investimentos"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 28

    row = 1
    if LOGO_PNG.exists():
        logo = XLImage(str(LOGO_PNG))
        logo.width = 220
        logo.height = 79
        ws.add_image(logo, "B1")
        row = 6
    else:
        ws.merge_cells("B1:D1")
        ws["B1"] = "3M AGROPECUÁRIA"
        ws["B1"].font = Font(name="Arial", size=18, bold=True, color=NAVY)
        ws["B1"].alignment = Alignment(horizontal="center")
        row = 3

    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    title_cell = ws.cell(row=row, column=2, value="INVESTIMENTOS")
    title_cell.font = Font(name="Arial", size=20, bold=True, color=NAVY)
    title_cell.alignment = Alignment(horizontal="center")
    row += 2

    for title, header_color, body_color in SECTIONS:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        header = ws.cell(row=row, column=2, value=title)
        header.font = Font(name="Arial", size=12, bold=True, color=WHITE)
        header.fill = PatternFill("solid", fgColor=header_color)
        header.alignment = Alignment(horizontal="center", vertical="center")
        header.border = thin_border()
        ws.row_dimensions[row].height = 28
        row += 1

        for col_idx, label in enumerate(COLUMNS, start=2):
            cell = ws.cell(row=row, column=col_idx, value=label)
            cell.font = Font(name="Arial", size=11, bold=True, color=header_color)
            cell.fill = PatternFill("solid", fgColor=body_color)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin_border()
        ws.row_dimensions[row].height = 22
        row += 1

        for i in range(INITIAL_ROWS):
            fill = BG_ALT if i % 2 else "FFFFFF"
            for col_idx in range(2, 5):
                cell = ws.cell(row=row, column=col_idx, value="")
                cell.font = Font(name="Arial", size=11, color="1E293B")
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border = thin_border()
            ws.row_dimensions[row].height = 24
            row += 1

        row += 1

    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
